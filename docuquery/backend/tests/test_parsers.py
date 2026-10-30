"""Unit tests for the multi-format parser registry, content validation,
zip-bomb guards, and the SSRF-guarded URL fetcher. No database required."""

import io
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from app.services import parsers
from app.services.url_fetch import UrlImportError, assert_public_url, fetch_document, suggest_name


def _write(tmp: Path, name: str, data: bytes) -> Path:
    p = tmp / name
    p.write_bytes(data)
    return p


def make_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("Contract Terms", level=1)
    d.add_paragraph("The seller shall deliver goods within thirty days.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Price"
    table.cell(1, 0).text = "Widget"
    table.cell(1, 1).text = "$5"
    d.save(str(path))


def make_pptx(path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly Review"
    slide.placeholders[1].text = "Revenue up twenty percent"
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])  # blank slide: skipped
    prs.save(str(path))


def make_xlsx(path: Path, rows: int = 2) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.append(["Invoice", "Amount"])
    for i in range(1, rows + 1):
        ws.append([f"INV-{i}", 100 + i])
    wb.save(str(path))


def make_pdf(path: Path, text: str = "Unique pdf probe content zqx.") -> None:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), text)
    doc.save(str(path))
    doc.close()


class TestParsers(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_txt_one_document_section(self):
        p = _write(self.tmp, "a.txt", "Hello world.\nSecond line.".encode())
        sections = parsers.parse_document(p, ".txt")
        self.assertEqual(len(sections), 1)
        self.assertEqual(sections[0].location, "Document")
        self.assertIn("Hello world.", sections[0].text)

    def test_md_sections_follow_headings(self):
        p = _write(self.tmp, "b.md", b"# Intro\nIntro body.\n\n## Details\nDetail text.")
        sections = parsers.parse_document(p, ".md")
        self.assertEqual([s.location for s in sections], ["Section: Intro", "Section: Details"])

    def test_md_frontmatter_skipped(self):
        p = _write(self.tmp, "c.md", b"---\ntitle: x\n---\n# Real\nBody.")
        sections = parsers.parse_document(p, ".md")
        self.assertEqual([s.location for s in sections], ["Section: Real"])

    def test_html_headings_and_tables(self):
        html = (
            "<html><head><title>T</title></head><body>"
            "<h1>Heading One</h1><p>Para one.</p>"
            "<table><tr><th>Col1</th><th>Col2</th></tr><tr><td>v1</td><td>v2</td></tr></table>"
            "</body></html>"
        )
        p = _write(self.tmp, "d.html", html.encode())
        sections = parsers.parse_document(p, ".html")
        self.assertEqual(sections[0].location, "Section: Heading One")
        self.assertIn("Para one.", sections[0].text)
        self.assertIn("Col1 | Col2", sections[0].text)
        self.assertIn("v1 | v2", sections[0].text)

    def test_docx_headings_and_tables(self):
        p = self.tmp / "e.docx"
        make_docx(p)
        sections = parsers.parse_document(p, ".docx")
        self.assertEqual(sections[0].location, "Section: Contract Terms")
        self.assertIn("thirty days", sections[0].text)
        self.assertIn("Item | Price", sections[0].text)
        self.assertIn("Widget | $5", sections[0].text)

    def test_pptx_slide_locations_skip_empty_slides(self):
        p = self.tmp / "f.pptx"
        make_pptx(p)
        sections = parsers.parse_document(p, ".pptx")
        self.assertEqual(len(sections), 1)
        self.assertEqual(sections[0].location, "Slide 1")
        self.assertIn("Revenue up twenty percent", sections[0].text)

    def test_xlsx_row_sections_repeat_header(self):
        p = self.tmp / "g.xlsx"
        make_xlsx(p, rows=2)
        with mock.patch.object(parsers, "get_settings") as fake:
            fake.return_value = mock.Mock(
                xlsx_rows_per_section=1, max_archive_entries=2000, max_uncompressed_mb=200
            )
            sections = parsers.parse_document(p, ".xlsx")
        self.assertEqual(len(sections), 2)
        self.assertEqual(sections[0].location, "Sheet 'Invoices' rows 2-2")
        self.assertEqual(sections[1].location, "Sheet 'Invoices' rows 3-3")
        self.assertIn("Columns: Invoice | Amount", sections[1].text)
        self.assertIn("Row 3: Invoice: INV-2 | Amount: 102", sections[1].text)

    def test_pdf_registry_path_unchanged(self):
        p = self.tmp / "h.pdf"
        make_pdf(p)
        sections = parsers.parse_document(p, ".pdf")
        self.assertEqual(len(sections), 1)
        self.assertEqual(sections[0].location, "Page 1")
        self.assertIn("Unique pdf probe content", sections[0].text)

    def test_unsupported_extension_rejected(self):
        p = _write(self.tmp, "i.exe", b"MZ binary")
        with self.assertRaisesRegex(ValueError, "Unsupported file type"):
            parsers.parse_document(p, ".exe")

    def test_magic_mismatch_rejected(self):
        p = _write(self.tmp, "fake.pdf", b"not a pdf at all")
        with self.assertRaisesRegex(ValueError, "PDF"):
            parsers.parse_document(p, ".pdf")

    def test_office_structure_mismatch_rejected(self):
        # A real DOCX claiming to be XLSX: zip signature matches but the
        # structural entries (word/) expose the forgery.
        docx_path = self.tmp / "real.docx"
        make_docx(docx_path)
        p = _write(self.tmp, "mislabeled.xlsx", docx_path.read_bytes())
        with self.assertRaisesRegex(ValueError, "XLSX"):
            parsers.parse_document(p, ".xlsx")

    def test_corrupt_office_file_rejected(self):
        p = _write(self.tmp, "corrupt.docx", b"PK\x03\x04" + b"garbage" * 100)
        with self.assertRaisesRegex(ValueError, "[Zz][Ii][Pp]"):
            parsers.parse_document(p, ".docx")

    def test_binary_text_file_rejected(self):
        p = _write(self.tmp, "bin.txt", b"\x00\x01\x02binary")
        with self.assertRaisesRegex(ValueError, "[Bb]inary"):
            parsers.parse_document(p, ".txt")

    def test_entry_count_zip_bomb_rejected(self):
        p = self.tmp / "bomb.docx"
        with zipfile.ZipFile(p, "w") as zf:
            zf.writestr("word/document.xml", "<w/>")
            for i in range(50):
                zf.writestr(f"junk/{i}", "x")
        with mock.patch.object(parsers, "get_settings") as fake:
            fake.return_value = mock.Mock(max_archive_entries=10, max_uncompressed_mb=200)
            with self.assertRaisesRegex(ValueError, "too many entries"):
                parsers.parse_document(p, ".docx")

    def test_uncompressed_size_zip_bomb_rejected(self):
        p = self.tmp / "big.docx"
        with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("word/document.xml", "<w/>")
            zf.writestr("big.bin", b"\0" * (5 * 1024 * 1024))  # compresses tiny
        with mock.patch.object(parsers, "get_settings") as fake:
            fake.return_value = mock.Mock(max_archive_entries=2000, max_uncompressed_mb=1)
            with self.assertRaisesRegex(ValueError, "safety limit"):
                parsers.parse_document(p, ".docx")

    def test_stored_extension_for_kind(self):
        self.assertEqual(parsers.stored_extension_for_kind("pdf"), ".pdf")
        self.assertEqual(parsers.stored_extension_for_kind("docx"), ".docx")
        self.assertEqual(parsers.stored_extension_for_kind("html"), ".html")
        self.assertEqual(parsers.stored_extension_for_kind("markdown"), ".md")
        self.assertEqual(parsers.stored_extension_for_kind("markdown", ".txt"), ".txt")
        self.assertEqual(parsers.stored_extension_for_kind("text", ".md"), ".md")
        self.assertEqual(parsers.stored_extension_for_kind("text"), ".txt")
        self.assertIsNone(parsers.stored_extension_for_kind("zip"))

    def test_sniff_kinds(self):
        self.assertEqual(parsers.sniff_kind_bytes(b"%PDF-1.7 ..."), "pdf")
        self.assertEqual(parsers.sniff_kind_bytes(b"<html><body>x</body></html>"), "html")
        self.assertEqual(parsers.sniff_kind_bytes(b"# Title\n\ntext"), "markdown")
        self.assertEqual(parsers.sniff_kind_bytes(b"just plain text"), "text")


class TestUrlFetch(unittest.TestCase):
    """SSRF guard + capped fetcher. IP-literal URLs avoid any DNS dependency."""

    def test_rejects_non_http_schemes(self):
        for url in ("ftp://example.com/x", "file:///etc/passwd", "gopher://x.io"):
            with self.assertRaises(UrlImportError):
                assert_public_url(url)

    def test_rejects_embedded_credentials(self):
        with self.assertRaises(UrlImportError):
            assert_public_url("https://user:pass@example.com/x")

    def test_rejects_loopback_and_private_targets(self):
        for url in (
            "http://127.0.0.1/x",
            "http://[::1]/x",
            "http://10.1.2.3/x",
            "http://192.168.1.1/x",
            "http://172.16.0.9/x",
            "http://169.254.169.254/latest/meta-data",  # cloud metadata endpoint
            "http://0.0.0.0/x",
        ):
            with self.assertRaises(UrlImportError, msg=url):
                assert_public_url(url)

    def test_rejects_unresolvable_hostname(self):
        with self.assertRaises(UrlImportError):
            assert_public_url("http://this-host-cannot-resolve-9x7q.invalid/")

    def test_accepts_public_ip_literal(self):
        self.assertEqual(assert_public_url("http://8.8.8.8/doc.pdf"), "http://8.8.8.8/doc.pdf")

    def _transport(self, handler) -> object:
        import httpx

        return httpx.MockTransport(handler)

    def test_fetch_document_returns_content_and_type(self):
        import httpx

        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(
                200,
                text="<html><body>Report body</body></html>",
                headers={"Content-Type": "text/html; charset=utf-8"},
            )

        content, ctype = fetch_document(
            "http://8.8.8.8/report.html",
            max_bytes=1024 * 1024,
            timeout_seconds=5,
            transport=self._transport(handler),
        )
        self.assertIn(b"Report body", content)
        self.assertEqual(ctype, "text/html")

    def test_fetch_document_rejects_disallowed_content_type(self):
        import httpx

        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(200, content=b"\xff\xd8\xff\xe0fakejpeg", headers={
                "Content-Type": "image/jpeg"
            })

        with self.assertRaises(UrlImportError):
            fetch_document(
                "http://8.8.8.8/pic.jpg",
                max_bytes=1024 * 1024,
                timeout_seconds=5,
                transport=self._transport(handler),
            )

    def test_fetch_document_enforces_size_cap(self):
        import httpx

        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(200, content=b"x" * 100)

        with self.assertRaisesRegex(UrlImportError, "limit"):
            fetch_document(
                "http://8.8.8.8/big.txt",
                max_bytes=10,
                timeout_seconds=5,
                transport=self._transport(handler),
            )

    def test_fetch_document_blocks_redirect_to_private_ip(self):
        import httpx

        def handler(request: httpx.Request) -> httpx.Response:
            if request.url.path == "/start":
                return httpx.Response(302, headers={"Location": "http://127.0.0.1/secret"})
            return httpx.Response(200, text="leak")

        with self.assertRaises(UrlImportError):
            fetch_document(
                "http://8.8.8.8/start",
                max_bytes=1024 * 1024,
                timeout_seconds=5,
                transport=self._transport(handler),
            )

    def test_fetch_document_follows_safe_redirects(self):
        import httpx

        def handler(request: httpx.Request) -> httpx.Response:
            if request.url.path == "/start":
                return httpx.Response(302, headers={"Location": "http://8.8.8.8/final.txt"})
            return httpx.Response(200, text="final content")

        content, _ = fetch_document(
            "http://8.8.8.8/start",
            max_bytes=1024 * 1024,
            timeout_seconds=5,
            transport=self._transport(handler),
        )
        self.assertEqual(content, b"final content")

    def test_suggest_name_uses_url_stem(self):
        self.assertEqual(suggest_name("https://ex.com/reports/annual.pdf", ".pdf"), "annual.pdf")
        self.assertEqual(suggest_name("https://ex.com/", ".html"), "ex.com.html")
        self.assertEqual(suggest_name("https://ex.com/no-ext", ".txt"), "no-ext.txt")


if __name__ == "__main__":
    unittest.main()
