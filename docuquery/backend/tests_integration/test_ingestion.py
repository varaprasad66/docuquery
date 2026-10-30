"""End-to-end ingestion + retrieval: upload -> pipeline -> ready -> searchable,
including hostile files and degraded modes."""

import io
from pathlib import Path
import uuid

from tests_integration import IntegrationTestCase, make_encrypted_pdf, make_pdf


class TestIngestionPipeline(IntegrationTestCase):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        self.email = f"ing-{uuid.uuid4().hex[:8]}@test.io"
        self.user = await self.register(self.email)
        self.token = self.user["access_token"]

    async def test_full_journey_upload_process_search(self):
        doc = await self.upload_pdf(
            self.token,
            filename="handbook.pdf",
            text="Employees receive twenty five days of annual leave per calendar year.",
        )
        self.assertEqual(doc["status"], "processing")
        await self.process_document(doc["id"])
        ready = await self.wait_ready(doc["id"], self.token)
        self.assertEqual(ready["status"], "ready")
        self.assertGreater(ready["chunk_count"], 0)
        self.assertEqual(ready["progress"], 1.0)

        r = await self.client.post(
            "/api/search",
            json={"query": "annual leave days per year"},
            headers=self.auth_headers(self.token),
        )
        self.assertEqual(r.status_code, 200)
        results = r.json()["results"]
        self.assertTrue(len(results) > 0, "Indexed content must be retrievable")
        self.assertIn("handbook.pdf", [x["document_name"] for x in results])

    async def test_reject_non_pdf(self):
        r = await self.client.post(
            "/api/documents",
            files={"file": ("fake.pdf", io.BytesIO(b"NOT A PDF"), "application/pdf")},
            headers=self.auth_headers(self.token),
        )
        self.assertEqual(r.status_code, 422)

    async def test_reject_empty_file(self):
        r = await self.client.post(
            "/api/documents",
            files={"file": ("empty.pdf", io.BytesIO(b""), "application/pdf")},
            headers=self.auth_headers(self.token),
        )
        self.assertEqual(r.status_code, 422)

    async def test_reject_pdf_forgery(self):
        # %PDF- magic header but garbage body: must fail processing, not crash.
        doc = await self.upload_pdf(
            self.token, filename="forged.pdf", content=b"%PDF-1.7 garbage not a real pdf"
        )
        await self.process_document(doc["id"])
        ready = await self.wait_ready(doc["id"], self.token)
        self.assertEqual(ready["status"], "failed")
        self.assertTrue(ready["error"])

    async def test_reject_encrypted_pdf(self):
        encrypted = make_encrypted_pdf()
        self.assertTrue(encrypted.startswith(b"%PDF-"))
        doc = await self.upload_pdf(self.token, filename="secret.pdf", content=encrypted)
        await self.process_document(doc["id"])
        ready = await self.wait_ready(doc["id"], self.token)
        self.assertEqual(ready["status"], "failed")
        self.assertIn("assword", ready["error"])  # readable message, not a stack trace

    async def test_delete_document_cascades_chunks_and_removes_from_search(self):
        doc = await self.upload_pdf(self.token, filename="deleteme.pdf", text="Zeppelin unique content xyzz y.")
        await self.process_document(doc["id"])
        doc_id = doc["id"]

        r = await self.client.post(
            "/api/search",
            json={"query": "Zeppelin unique content"},
            headers=self.auth_headers(self.token),
        )
        self.assertIn("deleteme.pdf", [x["document_name"] for x in r.json()["results"]])

        r = await self.client.delete(
            f"/api/documents/{doc_id}", headers=self.auth_headers(self.token)
        )
        self.assertEqual(r.status_code, 204)

        r = await self.client.post(
            "/api/search",
            json={"query": "Zeppelin unique content"},
            headers=self.auth_headers(self.token),
        )
        self.assertNotIn("deleteme.pdf", [x["document_name"] for x in r.json()["results"]])

    async def test_embedding_outage_degrades_to_bm25(self):
        from unittest.mock import patch

        doc = await self.upload_pdf(self.token, filename="bm25doc.pdf", text="Quantum flux capacitor settings.")
        await self.process_document(doc["id"])

        # Simulate the embedding provider being down at query time.
        async def boom(_texts):
            raise RuntimeError("embedding provider down")

        with patch("app.services.retrieval.embed_texts", side_effect=boom):
            r = await self.client.post(
                "/api/search",
                json={"query": "Quantum flux capacitor"},
                headers=self.auth_headers(self.token),
            )
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertIn("bm25", body["methods"])
        self.assertNotIn("vector", body["methods"])
        self.assertIn("bm25doc.pdf", [x["document_name"] for x in body["results"]])

    async def test_search_filter_by_document_ids(self):
        d1 = await self.upload_pdf(self.token, filename="one.pdf", text="Rooibos tea facts sheet unique alpha.")
        d2 = await self.upload_pdf(self.token, filename="two.pdf", text="Rooibos tea facts sheet unique beta.")
        await self.process_document(d1["id"])
        await self.process_document(d2["id"])

        r = await self.client.post(
            "/api/search",
            json={"query": "Rooibos tea facts", "document_ids": [d1["id"]]},
            headers=self.auth_headers(self.token),
        )
        names = {x["document_name"] for x in r.json()["results"]}
        if names:  # at least one hit must be respected the filter
            self.assertEqual(names, {"one.pdf"})


class TestMultiFormatIngestion(IntegrationTestCase):
    """Phase 1: DOCX/PPTX/XLSX/HTML/TXT/MD ingestion, location citations,
    and URL import."""

    async def asyncSetUp(self):
        await super().asyncSetUp()
        self.email = f"mf-{uuid.uuid4().hex[:8]}@test.io"
        self.user = await self.register(self.email)
        self.token = self.user["access_token"]

    async def upload_bytes(self, filename: str, content: bytes, mime: str) -> dict:
        r = await self.client.post(
            "/api/documents",
            files={"file": (filename, io.BytesIO(content), mime)},
            headers=self.auth_headers(self.token),
        )
        self.assertEqual(r.status_code, 201, r.text)
        return r.json()

    async def test_docx_upload_process_search_with_location(self):
        import docx as docx_lib
        import tempfile

        d = docx_lib.Document()
        d.add_heading("Leave Policy", level=1)
        d.add_paragraph("Sabbatical requests require sixty days notice unique marker qzx.")
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "policy.docx"
            d.save(str(p))
            doc = await self.upload_bytes("policy.docx", p.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        await self.process_document(doc["id"])
        ready = await self.wait_ready(doc["id"], self.token)
        self.assertEqual(ready["status"], "ready", ready.get("error"))

        r = await self.client.post(
            "/api/search",
            json={"query": "sabbatical sixty days notice"},
            headers=self.auth_headers(self.token),
        )
        results = r.json()["results"]
        self.assertTrue(results)
        hit = next(x for x in results if x["document_name"] == "policy.docx")
        self.assertEqual(hit["location"], "Section: Leave Policy")

    async def test_xlsx_header_repeated_and_row_locations(self):
        from openpyxl import Workbook
        import tempfile

        wb = Workbook()
        ws = wb.active
        ws.title = "Payroll"
        ws.append(["Employee", "Salary"])
        for i in range(1, 61):  # 60 rows -> two sections at 50 rows/section
            ws.append([f"Emp-{i}", 1000 + i])
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "payroll.xlsx"
            wb.save(str(p))
            doc = await self.upload_bytes("payroll.xlsx", p.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        await self.process_document(doc["id"])
        ready = await self.wait_ready(doc["id"], self.token)
        self.assertEqual(ready["status"], "ready", ready.get("error"))
        self.assertGreaterEqual(ready["page_count"], 2)

        r = await self.client.post(
            "/api/search",
            json={"query": "Emp-57"},
            headers=self.auth_headers(self.token),
        )
        results = r.json()["results"]
        self.assertTrue(results)
        hit = next(x for x in results if x["document_name"] == "payroll.xlsx")
        self.assertEqual(hit["location"], "Sheet 'Payroll' rows 52-61")
        self.assertIn("Columns: Employee | Salary", hit["snippet"])

    async def test_md_upload_search(self):
        md = b"# Onboarding\nYour first week checklist marker zzx begins here."
        doc = await self.upload_bytes("onboarding.md", md, "text/markdown")
        await self.process_document(doc["id"])
        ready = await self.wait_ready(doc["id"], self.token)
        self.assertEqual(ready["status"], "ready", ready.get("error"))

        r = await self.client.post(
            "/api/search",
            json={"query": "first week checklist"},
            headers=self.auth_headers(self.token),
        )
        hits = [x for x in r.json()["results"] if x["document_name"] == "onboarding.md"]
        self.assertTrue(hits)
        self.assertEqual(hits[0]["location"], "Section: Onboarding")

    async def test_reject_office_zip_bomb(self):
        import zipfile
        from unittest.mock import Mock, patch

        import app.api.documents as docs_api
        from app.services import parsers as parsers_svc

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("word/document.xml", "<w/>")
            zf.writestr("big.bin", b"\0" * (5 * 1024 * 1024))
        tiny = Mock(max_upload_mb=50, max_uncompressed_mb=1, max_archive_entries=2000)
        # The upload endpoint reads settings from its own module; the zip guard
        # reads them from the parsers module.
        with (
            patch.object(docs_api, "get_settings", return_value=tiny),
            patch.object(parsers_svc, "get_settings", return_value=tiny),
        ):
            r = await self.client.post(
                "/api/documents",
                files={"file": ("bomb.docx", io.BytesIO(buf.getvalue()), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                headers=self.auth_headers(self.token),
            )
        self.assertEqual(r.status_code, 422)
        self.assertIn("safety limit", r.json()["detail"])

    async def test_reject_renamed_extension(self):
        # DOCX bytes named .xlsx: signature check passes, structural check fails.
        import docx as docx_lib
        import tempfile

        d = docx_lib.Document()
        d.add_paragraph("Masquerading content.")
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "m.docx"
            d.save(str(p))
            r = await self.client.post(
                "/api/documents",
                files={"file": ("m.xlsx", io.BytesIO(p.read_bytes()), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                headers=self.auth_headers(self.token),
            )
        self.assertEqual(r.status_code, 422)
        self.assertIn("XLSX", r.json()["detail"])

    async def test_url_import_rejects_private_targets(self):
        for url in ("http://127.0.0.1/secret.pdf", "http://169.254.169.254/meta", "ftp://x/f.pdf"):
            r = await self.client.post(
                "/api/documents/url",
                json={"url": url},
                headers=self.auth_headers(self.token),
            )
            self.assertEqual(r.status_code, 422, url)

    async def test_url_import_happy_path_pdf(self):
        from unittest.mock import patch

        import app.api.documents as docs_api

        pdf_bytes = make_pdf(text="Imported via url unique marker plq.")

        def fake_fetch(url, *, max_bytes, timeout_seconds):
            return pdf_bytes, "application/pdf"

        with patch.object(docs_api, "fetch_document", side_effect=fake_fetch):
            r = await self.client.post(
                "/api/documents/url",
                json={"url": "http://8.8.8.8/reports/annual.pdf"},
                headers=self.auth_headers(self.token),
            )
        self.assertEqual(r.status_code, 201, r.text)
        doc = r.json()
        self.assertEqual(doc["original_name"], "annual.pdf")

        await self.process_document(doc["id"])
        ready = await self.wait_ready(doc["id"], self.token)
        self.assertEqual(ready["status"], "ready", ready.get("error"))
